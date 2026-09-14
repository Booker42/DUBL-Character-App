import importlib.util
import tempfile
import unittest
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / "tools" / "import_android_0_4_content.py"
spec = importlib.util.spec_from_file_location("import_android_0_4_content", SCRIPT)
module = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = module
assert spec.loader
spec.loader.exec_module(module)


class ArchmageImportTest(unittest.TestCase):
    def test_heading5_supported_spell_is_imported(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "archmage.docx"
            doc = Document()
            doc.add_paragraph("Клинок Короля-Ворона", style="Heading 5")
            doc.add_paragraph(
                "Школа: Боевая магия\n"
                "Стоимость: 4\n"
                "Время сотворения: 1 ОД\n"
                "Дальность: Оружие в руках\n"
                "Длительность: Мгновенно\n"
                "Описание: Призрачный клинок.\n"
                "Усиление: За 5 маны усиление."
            )
            doc.save(path)

            spells = module.parse_archmage_additions(path)
            self.assertEqual(["Клинок Короля-Ворона"], [spell["name"] for spell in spells])
            self.assertEqual("Боевая магия", spells[0]["school"])
            self.assertEqual(4, spells[0]["cost"])

    def test_heading5_unsupported_school_is_filtered(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "archmage.docx"
            doc = Document()
            doc.add_paragraph("Запрещённый спелл", style="Heading 5")
            doc.add_paragraph(
                "Школа: Элементалистика\n"
                "Стоимость: 4\n"
                "Описание: Не импортировать."
            )
            doc.save(path)

            self.assertEqual([], module.parse_archmage_additions(path))


if __name__ == "__main__":
    unittest.main()
