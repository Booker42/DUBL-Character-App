import hashlib
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from docx import Document

from tools.rulebook.extract_docx import extract_docx


def test_extract_docx_preserves_order_heading_paths_tables_and_hash(tmp_path: Path):
    source = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("Root", level=1)
    doc.add_paragraph("Alpha  text")
    doc.add_heading("Child", level=2)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "C"
    table.cell(1, 1).text = "D"
    doc.add_paragraph("Omega")
    doc.save(source)

    ir = extract_docx(source)

    assert ir["schemaVersion"] == 1
    assert ir["source"]["filename"] == "sample.docx"
    assert ir["source"]["sha256"] == hashlib.sha256(source.read_bytes()).hexdigest()
    assert [b["id"] for b in ir["blocks"]] == ["p-000001", "p-000002", "p-000003", "t-000004", "p-000005"]
    assert ir["blocks"][1]["text"] == "Alpha  text"
    assert ir["blocks"][1]["headingPath"] == ["Root"]
    assert ir["blocks"][2]["headingLevel"] == 2
    assert ir["blocks"][3]["headingPath"] == ["Root", "Child"]
    assert ir["blocks"][3]["rows"] == [["A", "B"], ["C", "D"]]
    assert ir["blocks"][4]["headingPath"] == ["Root", "Child"]
