from __future__ import annotations

from pathlib import Path
import sys

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from tools.rulebook.extract_docx import extract_docx


def _bindings() -> dict:
    return {
        "schemaVersion": 1,
        "bindings": [
            {"id": "athletics", "name": "Атлетика", "category": "PHYSICAL", "defaultAttributeHint": "STRENGTH", "template": False},
            {"id": "riding", "name": "Верховая езда", "category": "PHYSICAL", "defaultAttributeHint": "DEXTERITY", "template": False},
            {"id": "computers", "name": "Компьютеры", "category": "TECHNICAL", "defaultAttributeHint": "INTELLIGENCE", "template": False},
        ],
    }


def _source(tmp_path: Path) -> Path:
    source = tmp_path / "skills.docx"
    doc = Document()
    doc.add_heading("Умения", 1)
    table = doc.add_table(rows=1, cols=5)
    headers = ["Умение", "Описание", "Авто 6", "Авто 12", "Используется нетренированным"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for values in [
        ["Атлетика", "Прыжки, лазание, плавание", "Да", "Нет", "Да"],
        ["Верховая езда", "Держаться в седле", "Нет", "Нет", "Да (-2)"],
        ["Неизвестное умение", "Тест", "Нет", "Да", "Нет"],
    ]:
        row = table.add_row()
        for cell, value in zip(row.cells, values):
            cell.text = value
    doc.add_heading("Стоимость умений", 4)
    costs = doc.add_table(rows=1, cols=3)
    for cell, value in zip(costs.rows[0].cells, ["Значение", "Стоимость", "Стоимость для поднятия ранга с ноля"]):
        cell.text = value
    for values in [["0 → 1", "10", "10"], ["1 → 2", "20", "30"]]:
        row = costs.add_row()
        for cell, value in zip(row.cells, values):
            cell.text = value
    doc.save(source)
    return source


def test_import_skills_reads_canonical_table_costs_and_runtime_bindings(tmp_path: Path):
    from tools.rulebook.import_skills import import_skills

    raw = extract_docx(_source(tmp_path))
    payload, diagnostics = import_skills(raw, _bindings(), "core")

    assert payload["rankCosts"] == [0, 10, 30]
    assert [item["id"] for item in payload["skills"][:2]] == ["athletics", "riding"]
    athletics = payload["skills"][0]
    assert athletics == {
        "id": "athletics",
        "name": "Атлетика",
        "description": "Прыжки, лазание, плавание",
        "auto6": "Да",
        "auto12": "Нет",
        "untrained": "YES",
        "category": "PHYSICAL",
        "defaultAttributeHint": "STRENGTH",
        "template": False,
        "sourceRefs": ["core:t-000002"],
        "sourceRow": 1,
    }
    riding = payload["skills"][1]
    assert riding["untrained"] == "YES_MINUS_2"
    assert payload["rankCostSourceRefs"] == ["core:t-000004"]

    kinds = {(item["kind"], item["subject"], item["severity"]) for item in diagnostics}
    assert ("skill-runtime-binding-missing", "skills:Неизвестное умение", "warning") in kinds
    assert ("skill-binding-not-in-base-table", "skills:computers", "warning") in kinds


def test_import_skills_requires_exactly_one_base_skill_and_rank_cost_table(tmp_path: Path):
    from tools.rulebook.import_skills import import_skills

    source = tmp_path / "empty.docx"
    Document().save(source)
    raw = extract_docx(source)
    try:
        import_skills(raw, _bindings(), "core")
    except ValueError as exc:
        assert "base skill table" in str(exc)
    else:
        raise AssertionError("missing base skill table must fail structurally")


def test_import_skills_links_binding_only_skill_to_later_heading_but_keeps_unspecified_table_fields(tmp_path: Path):
    from tools.rulebook.import_skills import import_skills

    source = tmp_path / "skills_with_later_computers.docx"
    doc = Document()
    doc.add_heading("Умения", 1)
    table = doc.add_table(rows=1, cols=5)
    for cell, value in zip(table.rows[0].cells, ["Умение", "Описание", "Авто 6", "Авто 12", "Используется нетренированным"]):
        cell.text = value
    row = table.add_row()
    for cell, value in zip(row.cells, ["Атлетика", "Прыжки", "Да", "Нет", "Да"]):
        cell.text = value
    costs = doc.add_table(rows=1, cols=3)
    for cell, value in zip(costs.rows[0].cells, ["Значение", "Стоимость", "Стоимость для поднятия ранга с ноля"]):
        cell.text = value
    row = costs.add_row()
    for cell, value in zip(row.cells, ["0 → 1", "10", "10"]):
        cell.text = value
    doc.add_heading("Навыки", 1)
    doc.add_heading("Компьютеры", 3)
    doc.add_paragraph("Требование: Компьютеры 4")
    doc.add_paragraph("Проверка компьютеры+интеллект.")
    doc.save(source)

    raw = extract_docx(source)
    payload, diagnostics = import_skills(raw, _bindings(), "core")
    computers = next(item for item in payload["skills"] if item["id"] == "computers")
    assert computers["sourceKind"] == "supplemental-heading"
    assert computers["sourceRefs"] == ["core:p-000005"]
    assert computers["untrained"] == "UNSPECIFIED"
    assert computers["auto6"] == "Не указано в базовой таблице"
    assert computers["auto12"] == "Не указано в базовой таблице"
    assert any(d["kind"] == "skill-base-table-fields-missing" and d["subject"] == "skills:computers" for d in diagnostics)
    assert not any(d["kind"] == "skill-binding-not-in-base-table" and d["subject"] == "skills:computers" for d in diagnostics)
